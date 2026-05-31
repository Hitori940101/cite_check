"""Unit tests for the exporter module (CSV, Excel, BibTeX).

Covers all export functions with various data scenarios including
empty results, mixed statuses, Chinese characters, and file I/O.
"""

import csv
import io
import tempfile
from pathlib import Path

import pytest

from refchecker.core.exporter import (
    _count_sources,
    _escape_bibtex,
    _find_best_source,
    export_bibtex,
    export_bibtex_to_file,
    export_csv,
    export_csv_to_file,
    export_excel,
)
from refchecker.core.models import (
    AdapterMatch,
    ReferenceItem,
    VerificationResult,
    VerificationStatus,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _ref(**overrides: object) -> ReferenceItem:
    """Create a ReferenceItem with sensible defaults."""
    defaults = {
        "title": "Attention Is All You Need",
        "authors": ["Ashish Vaswani", "Noam Shazeer", "Niki Parmar"],
        "year": 2017,
        "journal": "NeurIPS",
        "doi": "10.48550/arXiv.1706.03762",
        "volume": "30",
        "entry_type": "article",
        "citation_key": "vaswani2017",
    }
    defaults.update(overrides)
    return ReferenceItem(**defaults)  # type: ignore[arg-type]


def _match(adapter: str = "crossref", score: float = 0.95, found: bool = True) -> AdapterMatch:
    """Create an AdapterMatch."""
    return AdapterMatch(
        adapter_name=adapter,
        found=found,
        composite_score=score,
        matched_title="Attention Is All You Need" if found else None,
    )


def _result(
    status: VerificationStatus = VerificationStatus.VERIFIED,
    best_score: float = 0.95,
    matches: list[AdapterMatch] | None = None,
    **ref_overrides: object,
) -> VerificationResult:
    """Create a VerificationResult with defaults."""
    if matches is None:
        matches = [_match(score=best_score)]
    return VerificationResult(
        reference=_ref(**ref_overrides),
        matches=matches,
        best_score=best_score,
        status=status,
    )


# ---------------------------------------------------------------------------
# _find_best_source
# ---------------------------------------------------------------------------


class TestFindBestSource:
    """Tests for _find_best_source helper."""

    def test_returns_matching_adapter(self) -> None:
        result = _result(matches=[_match("crossref", 0.90), _match("s2", 0.95)])
        assert _find_best_source(result) == "s2"

    def test_returns_first_on_tie(self) -> None:
        result = _result(
            matches=[_match("crossref", 0.90), _match("s2", 0.90)],
            best_score=0.90,
        )
        assert _find_best_source(result) == "crossref"

    def test_returns_empty_when_not_found(self) -> None:
        result = _result(
            matches=[AdapterMatch(adapter_name="crossref", found=False, composite_score=0.0)],
            best_score=0.0,
            status=VerificationStatus.LIKELY_FABRICATED,
        )
        assert _find_best_source(result) == ""

    def test_returns_empty_when_no_matches(self) -> None:
        result = VerificationResult(
            reference=_ref(),
            matches=[],
            best_score=0.0,
            status=VerificationStatus.LIKELY_FABRICATED,
        )
        assert _find_best_source(result) == ""


# ---------------------------------------------------------------------------
# _count_sources
# ---------------------------------------------------------------------------


class TestCountSources:
    """Tests for _count_sources helper."""

    def test_counts_adapters(self) -> None:
        result = _result(matches=[_match("crossref"), _match("s2")])
        assert _count_sources(result) == 2

    def test_zero_matches(self) -> None:
        result = VerificationResult(
            reference=_ref(),
            matches=[],
            best_score=0.0,
        )
        assert _count_sources(result) == 0


# ---------------------------------------------------------------------------
# _escape_bibtex
# ---------------------------------------------------------------------------


class TestEscapeBibtex:
    """Tests for _escape_bibtex helper."""

    def test_ampersand(self) -> None:
        assert _escape_bibtex("A & B") == r"A \& B"

    def test_percent(self) -> None:
        assert _escape_bibtex("100%") == r"100\%"

    def test_hash(self) -> None:
        assert _escape_bibtex("C#") == r"C\#"

    def test_no_special_chars(self) -> None:
        assert _escape_bibtex("Hello World") == "Hello World"

    def test_multiple_special(self) -> None:
        assert _escape_bibtex("A&B%C#") == r"A\&B\%C\#"


# ---------------------------------------------------------------------------
# CSV Export
# ---------------------------------------------------------------------------


class TestExportCsv:
    """Tests for export_csv and export_csv_to_file."""

    def test_basic_csv_output(self) -> None:
        results = [_result()]
        csv_str = export_csv(results)
        reader = csv.reader(io.StringIO(csv_str))
        rows = list(reader)

        # Header + 1 data row
        assert len(rows) == 2
        assert rows[0][0] == "#"
        assert rows[0][1] == "Title"
        assert rows[1][1] == "Attention Is All You Need"
        assert rows[1][6] == "verified"

    def test_csv_headers(self) -> None:
        csv_str = export_csv([])
        reader = csv.reader(io.StringIO(csv_str))
        headers = next(reader)
        expected = ["#", "Title", "Authors", "Year", "DOI", "Journal",
                     "Status", "Best Score", "Best Source", "Sources Checked"]
        assert headers == expected

    def test_empty_results(self) -> None:
        csv_str = export_csv([])
        reader = csv.StringIO(csv_str)
        lines = csv_str.strip().split("\n")
        # Only header row
        assert len(lines) == 1

    def test_multiple_results(self) -> None:
        results = [
            _result(status=VerificationStatus.VERIFIED, best_score=0.95),
            _result(
                status=VerificationStatus.SUSPICIOUS,
                best_score=0.70,
                title="Fake Paper",
            ),
        ]
        csv_str = export_csv(results)
        lines = csv_str.strip().split("\n")
        # Header + 2 data rows
        assert len(lines) == 3

    def test_chinese_characters(self) -> None:
        results = [_result(title="深度学习综述", authors=["张三", "李四"])]
        csv_str = export_csv(results)
        assert "深度学习综述" in csv_str
        assert "张三" in csv_str

    def test_missing_optional_fields(self) -> None:
        results = [_result(doi=None, journal=None, year=None)]
        csv_str = export_csv(results)
        reader = csv.reader(io.StringIO(csv_str))
        rows = list(reader)
        # DOI and Journal should be empty strings
        assert rows[1][3] == ""  # Year (None → "")
        assert rows[1][4] == ""  # DOI
        assert rows[1][5] == ""  # Journal

    def test_score_formatting(self) -> None:
        results = [_result(best_score=0.953)]
        csv_str = export_csv(results)
        assert "0.953" in csv_str

    def test_csv_file_write(self, tmp_path: Path) -> None:
        results = [_result()]
        path = tmp_path / "out.csv"
        export_csv_to_file(results, path)
        content = path.read_text(encoding="utf-8")
        assert "Attention Is All You Need" in content

    def test_csv_file_string_path(self, tmp_path: Path) -> None:
        results = [_result()]
        path = str(tmp_path / "out2.csv")
        export_csv_to_file(results, path)
        assert Path(path).exists()


# ---------------------------------------------------------------------------
# Excel Export
# ---------------------------------------------------------------------------


class TestExportExcel:
    """Tests for export_excel."""

    def test_basic_excel_export(self, tmp_path: Path) -> None:
        results = [_result()]
        path = tmp_path / "results.xlsx"
        export_excel(results, path)
        assert path.exists()
        assert path.stat().st_size > 0

    def test_excel_color_verified(self, tmp_path: Path) -> None:
        """Verified rows should have green background."""
        from openpyxl import load_workbook

        results = [_result(status=VerificationStatus.VERIFIED)]
        path = tmp_path / "green.xlsx"
        export_excel(results, path)

        wb = load_workbook(str(path))
        ws = wb.active
        # Row 2 (first data row), check fill color
        cell = ws.cell(row=2, column=1)
        assert cell.fill.start_color.rgb == "00C6EFCE"  # Green bg

    def test_excel_color_fabricated(self, tmp_path: Path) -> None:
        """Fabricated rows should have red background."""
        from openpyxl import load_workbook

        results = [_result(status=VerificationStatus.LIKELY_FABRICATED, best_score=0.3)]
        path = tmp_path / "red.xlsx"
        export_excel(results, path)

        wb = load_workbook(str(path))
        ws = wb.active
        cell = ws.cell(row=2, column=1)
        assert cell.fill.start_color.rgb == "00FFC7CE"  # Red bg

    def test_excel_color_suspicious(self, tmp_path: Path) -> None:
        """Suspicious rows should have yellow background."""
        from openpyxl import load_workbook

        results = [_result(status=VerificationStatus.SUSPICIOUS, best_score=0.70)]
        path = tmp_path / "yellow.xlsx"
        export_excel(results, path)

        wb = load_workbook(str(path))
        ws = wb.active
        cell = ws.cell(row=2, column=1)
        assert cell.fill.start_color.rgb == "00FFEB9C"  # Yellow bg

    def test_excel_color_unable_to_verify(self, tmp_path: Path) -> None:
        """Unable-to-verify rows should have blue/gray background."""
        from openpyxl import load_workbook

        results = [
            _result(
                status=VerificationStatus.UNABLE_TO_VERIFY,
                best_score=0.0,
                matches=[AdapterMatch(adapter_name="crossref", found=False, error="timeout")],
            )
        ]
        path = tmp_path / "blue.xlsx"
        export_excel(results, path)

        wb = load_workbook(str(path))
        ws = wb.active
        cell = ws.cell(row=2, column=1)
        assert cell.fill.start_color.rgb == "00D9E2F3"  # Blue/gray bg

    def test_excel_summary_statistics(self, tmp_path: Path) -> None:
        """Summary section should count statuses correctly."""
        from openpyxl import load_workbook

        results = [
            _result(status=VerificationStatus.VERIFIED),
            _result(status=VerificationStatus.VERIFIED),
            _result(status=VerificationStatus.SUSPICIOUS, best_score=0.70),
            _result(status=VerificationStatus.LIKELY_FABRICATED, best_score=0.3),
        ]
        path = tmp_path / "summary.xlsx"
        export_excel(results, path)

        wb = load_workbook(str(path))
        ws = wb.active
        # Summary starts at row len(results)+3 = 7
        assert ws.cell(row=7, column=1).value == "Summary"
        assert ws.cell(row=8, column=2).value == 2  # Verified
        assert ws.cell(row=9, column=2).value == 1  # Suspicious
        assert ws.cell(row=10, column=2).value == 1  # Fabricated
        assert ws.cell(row=11, column=2).value == 0  # Unable to verify

    def test_excel_empty_results(self, tmp_path: Path) -> None:
        """Empty results should produce valid Excel with headers only."""
        from openpyxl import load_workbook

        path = tmp_path / "empty.xlsx"
        export_excel([], path)

        wb = load_workbook(str(path))
        ws = wb.active
        # Only header row, summary starts at row 3
        assert ws.cell(row=1, column=1).value == "#"
        assert ws.cell(row=3, column=1).value == "Summary"

    def test_excel_chinese_content(self, tmp_path: Path) -> None:
        """Chinese characters should be preserved in Excel."""
        from openpyxl import load_workbook

        results = [_result(title="深度学习综述", authors=["张三"])]
        path = tmp_path / "chinese.xlsx"
        export_excel(results, path)

        wb = load_workbook(str(path))
        ws = wb.active
        assert ws.cell(row=2, column=2).value == "深度学习综述"

    def test_excel_multiple_results(self, tmp_path: Path) -> None:
        """Multiple results should produce correct number of data rows."""
        from openpyxl import load_workbook

        results = [
            _result(title=f"Paper {i}", best_score=0.9 + i * 0.01)
            for i in range(5)
        ]
        path = tmp_path / "multi.xlsx"
        export_excel(results, path)

        wb = load_workbook(str(path))
        ws = wb.active
        # Header + 5 data rows
        assert ws.max_row >= 6


# ---------------------------------------------------------------------------
# BibTeX Export
# ---------------------------------------------------------------------------


class TestExportBibtex:
    """Tests for export_bibtex and export_bibtex_to_file."""

    def test_basic_bibtex(self) -> None:
        results = [_result()]
        bib = export_bibtex(results)
        assert "@article{vaswani2017," in bib
        assert "title = {Attention Is All You Need}" in bib
        assert "author = {Ashish Vaswani and Noam Shazeer and Niki Parmar}" in bib
        assert "year = {2017}" in bib
        assert "doi = {10.48550/arXiv.1706.03762}" in bib

    def test_filters_fabricated_by_default(self) -> None:
        results = [
            _result(status=VerificationStatus.VERIFIED),
            _result(
                status=VerificationStatus.LIKELY_FABRICATED,
                best_score=0.3,
                citation_key="fake2020",
            ),
        ]
        bib = export_bibtex(results)
        assert "@article{vaswani2017," in bib
        assert "fake2020" not in bib

    def test_filters_by_include_statuses(self) -> None:
        results = [
            _result(status=VerificationStatus.VERIFIED),
            _result(
                status=VerificationStatus.SUSPICIOUS,
                best_score=0.70,
                citation_key="suspicious2020",
            ),
            _result(
                status=VerificationStatus.LIKELY_FABRICATED,
                best_score=0.3,
                citation_key="fake2020",
            ),
        ]
        # Only include fabricated
        bib = export_bibtex(results, include_statuses={VerificationStatus.LIKELY_FABRICATED})
        assert "fake2020" in bib
        assert "vaswani2017" not in bib
        assert "suspicious2020" not in bib

    def test_empty_results(self) -> None:
        bib = export_bibtex([])
        assert bib == ""

    def test_no_matching_status(self) -> None:
        results = [_result(status=VerificationStatus.LIKELY_FABRICATED, best_score=0.3)]
        bib = export_bibtex(results)  # Default: VERIFIED + SUSPICIOUS only
        assert bib == ""

    def test_entry_without_citation_key(self) -> None:
        results = [_result(citation_key=None)]
        bib = export_bibtex(results)
        # Should use auto-generated key starting with "ref_"
        assert "@article{ref_" in bib

    def test_entry_type_fallback(self) -> None:
        results = [_result(entry_type=None, citation_key="test1")]
        bib = export_bibtex(results)
        assert "@article{test1," in bib

    def test_optional_fields_present(self) -> None:
        results = [_result(volume="30", issue="1", pages="1-10", publisher="Springer")]
        bib = export_bibtex(results)
        assert "volume = {30}" in bib
        assert "number = {1}" in bib
        assert "pages = {1-10}" in bib
        assert "publisher = {Springer}" in bib

    def test_optional_fields_absent(self) -> None:
        results = [_ref(
            volume=None, issue=None, pages=None, publisher=None,
        )]
        result = VerificationResult(
            reference=_ref(volume=None, issue=None, pages=None, publisher=None),
            matches=[_match()],
            best_score=0.95,
            status=VerificationStatus.VERIFIED,
        )
        bib = export_bibtex([result])
        assert "volume" not in bib
        assert "number" not in bib
        assert "pages" not in bib
        assert "publisher" not in bib

    def test_bibtex_escape_in_title(self) -> None:
        results = [_result(title="R&D at 100% Efficiency")]
        bib = export_bibtex(results)
        assert r"R\&D at 100\% Efficiency" in bib

    def test_no_authors(self) -> None:
        ref = _ref(authors=[], citation_key="noauth")
        result = VerificationResult(
            reference=ref,
            matches=[_match()],
            best_score=0.95,
            status=VerificationStatus.VERIFIED,
        )
        bib = export_bibtex([result])
        assert "author" not in bib

    def test_no_year(self) -> None:
        ref = _ref(year=None, citation_key="noyear")
        result = VerificationResult(
            reference=ref,
            matches=[_match()],
            best_score=0.95,
            status=VerificationStatus.VERIFIED,
        )
        bib = export_bibtex([result])
        assert "year = {" not in bib

    def test_no_journal(self) -> None:
        ref = _ref(journal=None, citation_key="nojournal")
        result = VerificationResult(
            reference=ref,
            matches=[_match()],
            best_score=0.95,
            status=VerificationStatus.VERIFIED,
        )
        bib = export_bibtex([result])
        assert "journal = {" not in bib

    def test_bibtex_file_write(self, tmp_path: Path) -> None:
        results = [_result()]
        path = tmp_path / "out.bib"
        export_bibtex_to_file(results, path)
        content = path.read_text(encoding="utf-8")
        assert "@article{vaswani2017," in content

    def test_bibtex_file_with_include_statuses(self, tmp_path: Path) -> None:
        results = [
            _result(status=VerificationStatus.VERIFIED),
            _result(status=VerificationStatus.LIKELY_FABRICATED, best_score=0.3),
        ]
        path = tmp_path / "filtered.bib"
        export_bibtex_to_file(
            results, path,
            include_statuses={VerificationStatus.LIKELY_FABRICATED},
        )
        content = path.read_text(encoding="utf-8")
        assert "likely_fabricated" not in content  # status not in bib output

    def test_multiple_entries_separated(self) -> None:
        results = [
            _result(citation_key="a"),
            _result(status=VerificationStatus.SUSPICIOUS, best_score=0.70, citation_key="b"),
        ]
        bib = export_bibtex(results)
        # Two entries separated by double newline
        assert bib.count("@article") == 2
        parts = bib.split("\n\n")
        assert len(parts) == 2


# ---------------------------------------------------------------------------
# Round-trip: parse → export (integration-style but using public API)
# ---------------------------------------------------------------------------


class TestRoundTrip:
    """Verify exported data round-trips correctly."""

    def test_csv_round_trip_preserves_data(self) -> None:
        """CSV export preserves key reference data."""
        results = [
            _result(
                title="Deep Residual Learning",
                authors=["Kaiming He", "Xiangyu Zhang"],
                year=2016,
                doi="10.1109/CVPR.2016.90",
            )
        ]
        csv_str = export_csv(results)
        assert "Deep Residual Learning" in csv_str
        assert "Kaiming He" in csv_str
        assert "2016" in csv_str
        assert "10.1109/CVPR.2016.90" in csv_str
