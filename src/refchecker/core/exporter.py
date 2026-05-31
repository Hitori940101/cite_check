"""Export verification results to CSV, Excel, and clean BibTeX.

Supports:
- CSV: Plain-text spreadsheet with all verification fields
- Excel (.xlsx): Color-coded cells based on verification status
- BibTeX: Clean .bib with only verified/suspicious entries
"""

import csv
import io
from pathlib import Path
from typing import Sequence

from refchecker.core.logging import get_logger
from refchecker.core.models import VerificationResult, VerificationStatus

logger = get_logger(__name__)

# Status-to-color mapping for Excel
_STATUS_COLORS: dict[VerificationStatus, dict[str, str]] = {
    VerificationStatus.VERIFIED: {"bg": "C6EFCE", "fg": "006100"},
    VerificationStatus.SUSPICIOUS: {"bg": "FFEB9C", "fg": "9C6500"},
    VerificationStatus.LIKELY_FABRICATED: {"bg": "FFC7CE", "fg": "9C0006"},
    VerificationStatus.UNABLE_TO_VERIFY: {"bg": "D9E2F3", "fg": "1F4E79"},
    VerificationStatus.PENDING: {"bg": "FFFFFF", "fg": "000000"},
}

_CSV_HEADERS = [
    "#",
    "Title",
    "Authors",
    "Year",
    "DOI",
    "Journal",
    "Status",
    "Best Score",
    "Best Source",
    "Sources Checked",
]


def _find_best_source(result: VerificationResult) -> str:
    """Find the adapter name with the best score.

    Args:
        result: Verification result.

    Returns:
        Adapter name or empty string.
    """
    for match in result.matches:
        if match.found and match.composite_score == result.best_score:
            return match.adapter_name
    return ""


def _count_sources(result: VerificationResult) -> int:
    """Count the number of adapters that were queried.

    Args:
        result: Verification result.

    Returns:
        Number of adapters.
    """
    return len(result.matches)


# --- CSV Export ---


def export_csv(results: Sequence[VerificationResult]) -> str:
    """Export verification results to CSV string.

    Args:
        results: List of verification results.

    Returns:
        CSV string content.
    """
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_CSV_HEADERS)

    for i, r in enumerate(results, 1):
        ref = r.reference
        writer.writerow([
            i,
            ref.title,
            ref.display_authors,
            ref.year,
            ref.doi or "",
            ref.journal or "",
            r.status.value,
            f"{r.best_score:.3f}",
            _find_best_source(r),
            _count_sources(r),
        ])

    return buf.getvalue()


def export_csv_to_file(results: Sequence[VerificationResult], path: str | Path) -> None:
    """Export verification results to a CSV file.

    Args:
        results: List of verification results.
        path: Output file path.
    """
    content = export_csv(results)
    Path(path).write_text(content, encoding="utf-8")
    logger.info("export_csv", path=str(path), count=len(results))


# --- Excel Export ---


def export_excel(results: Sequence[VerificationResult], path: str | Path) -> None:
    """Export verification results to a color-coded Excel file.

    Args:
        results: List of verification results.
        path: Output file path (must end in .xlsx).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Verification Results"

    # Header row
    header_font = Font(bold=True)
    for col, header in enumerate(_CSV_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font

    # Data rows
    for i, r in enumerate(results, 2):
        ref = r.reference
        row_data = [
            i - 1,
            ref.title,
            ref.display_authors,
            ref.year,
            ref.doi or "",
            ref.journal or "",
            r.status.value,
            round(r.best_score, 3),
            _find_best_source(r),
            _count_sources(r),
        ]

        colors = _STATUS_COLORS.get(r.status, _STATUS_COLORS[VerificationStatus.PENDING])
        fill = PatternFill(start_color=colors["bg"], end_color=colors["bg"], fill_type="solid")
        font = Font(color=colors["fg"])

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.fill = fill
            cell.font = font

    # Auto-width columns
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    # Summary row
    summary_row = len(results) + 3
    ws.cell(row=summary_row, column=1, value="Summary").font = Font(bold=True)
    ws.cell(row=summary_row + 1, column=1, value="Verified")
    ws.cell(row=summary_row + 1, column=2, value=sum(1 for r in results if r.status == VerificationStatus.VERIFIED))
    ws.cell(row=summary_row + 2, column=1, value="Suspicious")
    ws.cell(row=summary_row + 2, column=2, value=sum(1 for r in results if r.status == VerificationStatus.SUSPICIOUS))
    ws.cell(row=summary_row + 3, column=1, value="Likely Fabricated")
    ws.cell(row=summary_row + 3, column=2, value=sum(1 for r in results if r.status == VerificationStatus.LIKELY_FABRICATED))
    ws.cell(row=summary_row + 4, column=1, value="Unable to Verify")
    ws.cell(row=summary_row + 4, column=2, value=sum(1 for r in results if r.status == VerificationStatus.UNABLE_TO_VERIFY))

    wb.save(str(path))
    logger.info("export_excel", path=str(path), count=len(results))


# --- BibTeX Export ---


def _escape_bibtex(text: str) -> str:
    """Escape special BibTeX characters.

    Args:
        text: Raw text.

    Returns:
        Escaped text.
    """
    return text.replace("&", r"\&").replace("%", r"\%").replace("#", r"\#")


def export_bibtex(
    results: Sequence[VerificationResult],
    *,
    include_statuses: set[VerificationStatus] | None = None,
) -> str:
    """Export verification results to clean BibTeX format.

    Only exports entries with the specified statuses. Defaults to
    verified and suspicious only.

    Args:
        results: List of verification results.
        include_statuses: Set of statuses to include. Default: VERIFIED + SUSPICIOUS.

    Returns:
        BibTeX string content.
    """
    if include_statuses is None:
        include_statuses = {VerificationStatus.VERIFIED, VerificationStatus.SUSPICIOUS}

    entries: list[str] = []
    for r in results:
        if r.status not in include_statuses:
            continue

        ref = r.reference
        entry_type = ref.entry_type or "article"
        key = ref.citation_key or f"ref_{hash(ref.title) % 100000}"

        lines = [f"@{entry_type}{{{key},"]
        lines.append(f"  title = {{{_escape_bibtex(ref.title)}}},")

        if ref.authors:
            author_str = " and ".join(ref.authors)
            lines.append(f"  author = {{{_escape_bibtex(author_str)}}},")

        if ref.year is not None:
            lines.append(f"  year = {{{ref.year}}},")

        if ref.journal:
            lines.append(f"  journal = {{{_escape_bibtex(ref.journal)}}},")

        if ref.volume:
            lines.append(f"  volume = {{{ref.volume}}},")

        if ref.issue:
            lines.append(f"  number = {{{ref.issue}}},")

        if ref.pages:
            lines.append(f"  pages = {{{ref.pages}}},")

        if ref.doi:
            lines.append(f"  doi = {{{ref.doi}}},")

        if ref.publisher:
            lines.append(f"  publisher = {{{_escape_bibtex(ref.publisher)}}},")

        # Remove trailing comma from last field
        if lines[-1].endswith(","):
            lines[-1] = lines[-1][:-1]

        lines.append("}")
        entries.append("\n".join(lines))

    return "\n\n".join(entries)


def export_bibtex_to_file(
    results: Sequence[VerificationResult],
    path: str | Path,
    *,
    include_statuses: set[VerificationStatus] | None = None,
) -> None:
    """Export verification results to a clean BibTeX file.

    Args:
        results: List of verification results.
        path: Output file path.
        include_statuses: Set of statuses to include.
    """
    content = export_bibtex(results, include_statuses=include_statuses)
    Path(path).write_text(content, encoding="utf-8")
    logger.info("export_bibtex", path=str(path), count=len(results))
